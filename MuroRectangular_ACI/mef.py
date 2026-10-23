"""MEF del muro y lectura de ETABS.

QUÉ HACE
  1. Lee de ETABS las fuerzas del pier (P, V2, V3, M2, M3) en la base y el tope
     de cada entrepiso, para cada combinación.
  2. Con esas fuerzas calcula las cargas que entran en cada losa (fuerza
     horizontal H, axial N y par C) y el peso propio de cada entrepiso W:
         H_i = V_i − V_(i+1)
         N_i = P_tope,i − P_(i+1)
         C_i = (M_i − M_(i+1)) + h_i·V_i
         W_i = P_i − P_tope,i
     (P_i, V_i, M_i: base del entrepiso i; P_tope,i: su tope.)
  3. Arma un modelo de elementos finitos del muro en su plano (tensión plana,
     elementos Q4 de 4 nudos), empotrado en la base, y le aplica H, N y C en
     los nudos de cada losa y W repartido en los planos interiores del
     entrepiso, como el peso del muro.
  4. Obtiene la fuerza de sección (P, V, M) en cada plano de nudos sumando las
     fuerzas nodales internas de los elementos que apoyan en ese plano. Es lo
     mismo que hace ETABS para dar sus fuerzas de pier.

POR QUÉ EL MEF DEVUELVE EXACTAMENTE LO DE ETABS
  El muro aislado, empotrado solo en la base, es ISOSTÁTICO: sus fuerzas de
  sección salen de la estática, no de la rigidez. Cargado con las fuerzas de
  ETABS, tiene que devolverlas exactas en la base de cada entrepiso, y entre
  losas V es constante, M es una recta de pendiente V (dM/dz = V) y P solo
  cambia por el peso propio.
  El contraste con ETABS comprueba, por tanto, que las cargas se transfirieron
  bien y que la suma de fuerzas nodales está bien hecha; en el tope, que la
  parte estática de ETABS cumple M_tope = M_base + h·V. Los desplazamientos
  del MEF no se usan: el δu del elemento de borde es dato (archivo de datos). Por lo
  mismo, en los tipos II y III basta modelar el alma: las alas cambiarían la
  rigidez, no las fuerzas de sección.

EL PESO PROPIO ENTRE LOSAS. En ETABS P baja del pie al tope de cada entrepiso
  por el peso del muro. Con AXIAL_CON_PESO_PROPIO = True (archivo de datos; lo
  normal) el MEF lo carga repartido en la altura, W_i = P_i − P_tope,i (la
  diferencia que da ETABS, con el factor de cada combinación), y devuelve el P
  de ETABS en la base Y en el tope; la demanda lleva P en recta entre los dos.
  Con False las cargas salen solo de las bases: el peso queda en la losa de
  arriba y P es constante en el entrepiso, con el valor de su base. V y M no
  cambian: el peso, centrado, no da cortante ni momento.

LOS EJES. ETABS da las fuerzas del pier en sus ejes locales 2 y 3, y cada pier
  tiene los suyos (AxisAngle). Este programa trabaja siempre en los mismos:
  X a lo largo del muro, del extremo izquierdo al derecho, e Y a través del
  espesor, hacia el ala. EJE_X y EJE_Y (archivo de datos) dicen a qué eje de
  ETABS apunta cada uno, y `a_ejes_del_programa` traduce las fuerzas.

  LA CONVENCIÓN DE ETABS, deducida de los propios datos con los casos
  estáticos (ver `comprobar_convencion`): V2 y V3 son positivos en el sentido
  de +2 y +3 (la fuerza que el tramo de arriba le hace al de abajo), M3 > 0
  comprime la cara +2 y M2 > 0 comprime la cara +3. En este programa el momento
  es NEGATIVO cuando comprime la cara +X (o +Y), así que, con X = +2:
      Vx = V2      Mx = −M3
  y con X = −2: Vx = −V2, Mx = +M3; igual con Y y el eje 3 (o al revés). Con
  esa regla dM/dz = V en las dos direcciones.

LA DIRECCIÓN TRANSVERSAL (fuera del plano) no se puede cargar en un modelo de
tensión plana. No hace falta: entre losas no entra carga, así que la parte
estática es
    My(z) = My_base + Vy·(z − z_base)
con los valores de la base de ETABS (comprobado contra su valor Top); la de
sismo de las envolventes, como la de Mx (ver LAS ENVOLVENTES).

LAS ENVOLVENTES. Las combinaciones con sismo (espectro) vienen de ETABS como
  filas Max y Min, y cada componente es su propio máximo: no son estados de un
  mismo instante. `Demanda.en` separa la parte de gravedad (la media) y la de
  sismo (la semidiferencia) y, con PERMUTAR_SIGNOS = True (archivo de datos),
  arma los 8 estados con los dos signos de la de sismo en P, Mx y My, como
  hace ETABS en su diseño ("permutaciones de máximos y mínimos"). V va con el
  signo de Mx y Vy con el de My, como en las filas de ETABS (en un voladizo,
  el cortante y el momento de la base van juntos).

  EL CORTE (§18.10.3.3.2) amplifica solo la parte de sismo, VuEh. Cada estado
  guarda su Vg y su VEh, que salen por separado de ETABS: de cada combinación
  con sismo, su combinación de gravedad (0.9D, 1.25(L+D)) y su caso de
  espectro de diseño (SEPARACION_DEL_CORTE, en el archivo de datos);
  `leer_separacion_del_corte` comprueba que la combinación de ETABS es
  gravedad ± sismo. La gravitacional no tiene sismo y no se amplifica. Esta
  separación es solo para el corte: lo demás usa las combinaciones tal como
  vienen.

  LA PARTE DE GRAVEDAD es un estado estático: sale del MEF (y en el tope
  coincide con ETABS). LA PARTE DE SISMO no: es la combinación modal (CQC) de
  los modos, |S| = √(Σ ρij·Si·Sj), y no cumple la estática; por eso cargar el
  MEF con la fila de la base no da el tope de ETABS (puede salir hasta con el
  signo cambiado).

  Dentro de un entrepiso el cortante de cada modo es constante y su momento,
  una recta: M_k(z) = a_k + z·b_k. SRSS y CQC son formas cuadráticas,
  S² = Σ ρij·Si·Sj, así que S(z)² = |a|² + 2z·<a,b> + z²·|b|², y los tres
  términos salen de lo que da ETABS: |a| = S_base, |b| = S_V (el cortante de
  sismo) y, del tope, 2h·<a,b> = S_tope² − S_base² − h²·S_V². Queda, con
  ζ = (z − z_base)/h:
      S(z)² = (1 − ζ)·S_base² + ζ·S_tope² − ζ·(1 − ζ)·h²·S_V²
  EXACTA para SRSS y CQC (comprobada contra una CQC con modos al azar), y da
  S_base y S_tope en los extremos. Con SISMO_ENTRE_BASE_Y_TOPE = True
  (archivo de datos; lo normal) la parte de sismo de M y My sale de ahí; la de
  P, V y Vy es constante en el entrepiso. La estática desde la base (False)
  da |S_base − z·S_V|, que es la cota INFERIOR: coincide
  con la exacta donde domina un modo y la subestima cerca de donde el momento
  de ese modo cambia de signo (ahí puede dar casi cero donde la exacta no lo es).
  Si el tope de ETABS no cae entre |S_base − h·S_V| y S_base + h·S_V (una
  combinación direccional ABS, p. ej., que no es cuadrática), se usa la recta
  entre base y tope, que acota por arriba cualquier combinación convexa.
  (Supone que el muro no tiene masas entre losas en ETABS; si está dividido en
  altura, sus masas intermedias curvan un poco el momento de cada modo.)

CARGAS DENTRO DEL ENTREPISO. Si en ETABS entra carga entre dos losas (un nivel
  intermedio, vigas a media altura…), la parte estática del tope no es la de
  la base: V cambia y M_tope ≠ M_base + V·h (`cargas_dentro_del_entrepiso`,
  más que el redondeo de ETABS). Como no se sabe dónde entra, la demanda de
  ese entrepiso es la ENVOLVENTE de dos lecturas: la estática desde la base
  (la del MEF) y la estática desde el tope de ETABS (los estados «(tope)»).
  Criterio del proyecto; el informe lo avisa («OJO») en el entrepiso donde
  pasa.

EL CENTROIDE. ETABS da los momentos respecto del centro de gravedad de su
  modelo de ejes (cada muro, un rectángulo sobre su línea media); la capacidad
  se calcula respecto del centroide de la sección real. Con alas no coinciden,
  y la demanda se traslada: M_real = M_ETABS + P·(x_ETABS − x_real).

UNIDADES: kN, m. Los materiales llegan en MPa y se pasan a kN/m² (× 1000).
"""

from __future__ import annotations

import math
from dataclasses import dataclass
from pathlib import Path

import numpy as np

import datos as D

MPA = 1000.0          # 1 MPa = 1000 kN/m²


class DatoInvalido(ValueError):
    """Un dato del archivo de datos no tiene sentido; el mensaje dice cuál y por qué."""


# =============================================================================
# 1. LECTURA DE ETABS
# =============================================================================


A_KN = {"tonf": 9.80665, "kN": 1.0}             # unidades de fuerza de ETABS → kN
A_KN_M = {"tonf-m": 9.80665, "kN-m": 1.0}       # unidades de momento de ETABS → kN·m
COMPONENTES = ("P", "V2", "V3", "M2", "M3")


def leer_fuerzas_etabs(casos=None) -> dict:
    """Fuerzas del pier, en kN y kN·m, con el signo y los ejes de ETABS.

    Devuelve un diccionario:
        fuerzas[(caso, paso)][piso][ubicación] = {"P":…, "V2":…, …}
    con piso = 0, 1, 2 (de abajo arriba), ubicación = "Bottom" o "Top" y
    paso = "Max", "Min" o "" (la gravitacional no tiene paso). Con `casos`, lee
    esos en vez de COMBINACIONES (los estáticos, que traen varios pasos
    "Step By Step": el paso se guarda con su número). La tabla puede venir en
    tonf y tonf-m o en kN y kN-m: se lee de su fila de unidades y se pasa a kN.
    """
    return _leer_hoja(set(D.COMBINACIONES if casos is None else casos), casos)


NECESARIAS = ("Story", "Pier", "Output Case", "Step Type", "Location", *COMPONENTES)
_FILAS: dict = {}          # (archivo, hoja, fecha del archivo) → filas: el Excel se abre una sola vez


def carpeta_de_los_datos() -> Path:
    """La carpeta del archivo de datos, la de su edificación: ahí están su
    etabs/ y su resultados/."""
    return Path(D.__file__).resolve().parent


def filas_de_la_hoja() -> list:
    """Las filas de la hoja HOJA_FUERZAS, tal como vienen (se leen una vez)."""
    ruta = carpeta_de_los_datos() / D.ARCHIVO_FUERZAS
    try:
        clave = (str(ruta), D.HOJA_FUERZAS, ruta.stat().st_mtime)
    except OSError:
        clave = None
    if clave in _FILAS:
        return _FILAS[clave]
    from openpyxl import load_workbook  # se importa aquí: solo hace falta para leer

    try:
        libro = load_workbook(ruta, read_only=True, data_only=True)
    except Exception as error:           # no existe, no es un .xlsx, está dañado…
        raise DatoInvalido(f"no se pudo abrir {D.ARCHIVO_FUERZAS} como libro de Excel "
                           f"({type(error).__name__}: {error})")
    try:
        if D.HOJA_FUERZAS not in libro.sheetnames:
            raise DatoInvalido(f"{D.ARCHIVO_FUERZAS} no tiene la hoja {D.HOJA_FUERZAS!r} (HOJA_FUERZAS); "
                               f"tiene {libro.sheetnames}")
        filas = list(libro[D.HOJA_FUERZAS].iter_rows(values_only=True))
    finally:
        libro.close()                    # en modo de solo lectura el archivo queda abierto
    if clave is not None:
        _FILAS[clave] = filas
    return filas


def _leer_hoja(nombres: set, casos) -> dict:
    """El cuerpo de `leer_fuerzas_etabs`, con las filas de la hoja."""
    filas = iter(filas_de_la_hoja())
    try:
        next(filas)                      # fila de título: "TABLE: Pier Forces"
        cabecera = list(next(filas))
        unidades = list(next(filas))
    except StopIteration:
        raise DatoInvalido(f"la hoja {D.HOJA_FUERZAS!r} de {D.ARCHIVO_FUERZAS} no trae el título, la "
                           f"cabecera y las unidades de la tabla Pier Forces de ETABS")
    col = {nombre: i for i, nombre in enumerate(cabecera)}
    faltan = [c for c in NECESARIAS if c not in col]
    if faltan:
        raise DatoInvalido(f"a la tabla de {D.ARCHIVO_FUERZAS} le faltan las columnas {faltan} (hace "
                           f"falta la tabla Pier Forces de ETABS tal como se exporta)")
    factor = {}
    for k in COMPONENTES:
        tabla = A_KN if k in ("P", "V2", "V3") else A_KN_M
        u = unidades[col[k]]
        if u not in tabla:
            raise DatoInvalido(f"la columna {k} de {D.ARCHIVO_FUERZAS} viene en {u!r}; este programa "
                               f"lee {' o '.join(repr(x) for x in tabla)}")
        factor[k] = tabla[u]

    fuerzas: dict = {}
    for f in filas:
        if f[col["Pier"]] != D.MURO or f[col["Output Case"]] not in nombres:
            continue
        if f[col["Story"]] not in D.PISOS:
            continue
        paso = f[col["Step Type"]] or ""
        if paso == "Step By Step":
            if "Step Number" not in col or f[col["Step Number"]] is None:
                raise DatoInvalido(f"{D.ARCHIVO_FUERZAS}: {f[col['Output Case']]} viene paso a paso pero "
                                   f"la tabla no trae la columna «Step Number»")
            paso = f"paso {int(f[col['Step Number']])}"
        clave = (f[col["Output Case"]], paso)
        piso = D.PISOS.index(f[col["Story"]])
        donde = f[col["Location"]]
        try:
            valores = {k: float(f[col[k]]) * factor[k] for k in COMPONENTES}
        except (TypeError, ValueError):
            raise DatoInvalido(f"en {D.ARCHIVO_FUERZAS}, la fila de {clave[0]} {clave[1]} "
                               f"{D.PISOS[piso]} {donde} tiene una fuerza que no es un número")
        por_piso = fuerzas.setdefault(clave, {}).setdefault(piso, {})
        if donde in por_piso:
            raise DatoInvalido(f"fila repetida en {D.ARCHIVO_FUERZAS}: {clave[0]} {clave[1]} "
                               f"{D.PISOS[piso]} {donde}")
        por_piso[donde] = valores

    for clave, por_piso in fuerzas.items():
        for piso in range(len(D.PISOS)):
            if set(por_piso.get(piso, {})) != {"Bottom", "Top"}:
                raise DatoInvalido(f"en {D.ARCHIVO_FUERZAS}, a {clave[0]} {clave[1]} le falta la base o "
                                   f"el tope del piso {D.PISOS[piso]} (¿PISOS coincide con ETABS?)")
    if not fuerzas and casos is None:
        raise DatoInvalido(f"no se encontró el pier {D.MURO} en {D.ARCHIVO_FUERZAS}")
    if casos is None:
        faltan = [c for c in D.COMBINACIONES if not any(k[0] == c for k in fuerzas)]
        if faltan:
            raise DatoInvalido(f"en {D.ARCHIVO_FUERZAS} no está, para el pier {D.MURO}: {faltan} (en "
                               f"COMBINACIONES van los nombres de ETABS, su «Output Case»)")
    return dict(sorted(fuerzas.items()))


# =============================================================================
# 1b. LOS EJES: DE ETABS A ESTE PROGRAMA
# =============================================================================


def ejes() -> tuple[tuple[str, float], tuple[str, float]]:
    """(eje de ETABS, signo) de X y de Y, leídos de EJE_X y EJE_Y."""
    salida = []
    for nombre in ("EJE_X", "EJE_Y"):
        valor = str(getattr(D, nombre)).strip().replace(" ", "")
        if valor not in ("+2", "-2", "+3", "-3", "2", "3"):
            raise DatoInvalido(f"{nombre} es \"+2\", \"-2\", \"+3\" o \"-3\", no {valor!r}")
        salida.append((valor[-1], -1.0 if valor.startswith("-") else 1.0))
    if salida[0][0] == salida[1][0]:
        raise DatoInvalido("EJE_X y EJE_Y tienen que ir a ejes distintos de ETABS (uno al 2 y "
                           "otro al 3)")
    return salida[0], salida[1]


def a_ejes_del_programa(registro: dict) -> dict:
    """{P, V2, V3, M2, M3} de ETABS → {P, Vx, Mx, Vy, My} de este programa.

    V va con el signo del eje; M, con signo contrario al momento que comprime
    la cara positiva de ese eje (M3 comprime +2, M2 comprime +3).
    """
    (ex, sx), (ey, sy) = ejes()
    V = {"2": registro["V2"], "3": registro["V3"]}
    M = {"2": registro["M3"], "3": registro["M2"]}     # el que comprime la cara +2 / +3
    return {"P": registro["P"], "Vx": sx * V[ex], "Mx": -sx * M[ex],
            "Vy": sy * V[ey], "My": -sy * M[ey]}


def mapear(fuerzas: dict) -> dict:
    """Las fuerzas de `leer_fuerzas_etabs`, en los ejes del programa."""
    return {clave: {s: {donde: a_ejes_del_programa(r) for donde, r in por_piso.items()}
                    for s, por_piso in pisos.items()}
            for clave, pisos in fuerzas.items()}


def direcciones_globales() -> dict:
    """Dirección en planta (global) de +2, +3, +X y +Y, con AXIS_ANGLE."""
    th = math.radians(D.AXIS_ANGLE)
    e = {"2": np.array([math.cos(th), math.sin(th)]),
         "3": np.array([-math.sin(th), math.cos(th)])}     # 3 = 2 girado +90°
    (ex, sx), (ey, sy) = ejes()
    return {"2": e["2"], "3": e["3"], "X": sx * e[ex], "Y": sy * e[ey]}


def nombre_global(v: np.ndarray) -> str:
    """'+X global', '−Y global' o el ángulo, para un vector unitario en planta."""
    for nombre, u in (("+X", (1, 0)), ("−X", (-1, 0)), ("+Y", (0, 1)), ("−Y", (0, -1))):
        if abs(v[0] - u[0]) < 1e-6 and abs(v[1] - u[1]) < 1e-6:
            return f"{nombre} global"
    return f"{math.degrees(math.atan2(v[1], v[0])):.1f}° desde +X global"


def comprobar_convencion() -> list[dict]:
    """Comprueba la convención de signos de ETABS con los casos estáticos.

    Un caso estático empuja en +X o +Y global. Si la convención es la que usa
    este programa, en la base del primer piso:
      - el cortante (V2·e2 + V3·e3) apunta hacia donde empuja la carga;
      - la cara comprimida por el momento (M3·e2 + M2·e3) también, como en un
        voladizo empujado por arriba.
    Devuelve una fila por caso y paso. Vacío si no hay casos estáticos.
    """
    casos = getattr(D, "CASOS_ESTATICOS", None) or {}
    if not casos:
        return []
    fuerzas = leer_fuerzas_etabs(casos.values())
    d = direcciones_globales()
    filas = []
    for clave, por_piso in fuerzas.items():
        sentido = next(k for k, v in casos.items() if v == clave[0])
        empuje = np.array([1.0, 0.0]) if sentido == "X" else np.array([0.0, 1.0])
        r = por_piso[0]["Bottom"]
        cortante = r["V2"] * d["2"] + r["V3"] * d["3"]
        comprimida = r["M3"] * d["2"] + r["M2"] * d["3"]
        filas.append({"caso": clave[0], "paso": clave[1], "registro": r,
                      "cortante": float(cortante @ empuje), "momento": float(comprimida @ empuje),
                      "cumple": float(cortante @ empuje) > 0 and float(comprimida @ empuje) > 0})
    return filas


# =============================================================================
# 2. LA MALLA
# =============================================================================


@dataclass
class Malla:
    xs: np.ndarray        # abscisas de los nudos (a lo largo del muro)
    zs: np.ndarray        # cotas de los planos de nudos
    losas: list           # índice del plano de cada losa (tope de cada entrepiso)
    bases: list           # índice del plano de la base de cada entrepiso

    @property
    def nx(self) -> int:  # nudos por plano
        return len(self.xs)

    def nudo(self, i: int, j: int) -> int:
        """Número del nudo en la columna i (x) y el plano j (z)."""
        return j * self.nx + i


def crear_malla() -> Malla:
    """Elementos de MALLA_MEF como máximo; cada entrepiso se divide en partes iguales."""
    n_x = math.ceil(D.LW / D.MALLA_MEF - 1e-9)
    xs = np.array([D.LW * i / n_x for i in range(n_x + 1)])
    zs = [0.0]
    bases, losas = [], []
    z0 = 0.0
    for h in D.ALTURAS:
        bases.append(len(zs) - 1)
        # al menos 2 elementos: el peso propio va en los planos interiores
        n_z = max(2, math.ceil(h / D.MALLA_MEF - 1e-9))
        zs.extend(z0 + h * j / n_z for j in range(1, n_z + 1))
        z0 += h
        losas.append(len(zs) - 1)
    return Malla(xs, np.array(zs), losas, bases)


# =============================================================================
# 3. EL ELEMENTO: Q4 DE TENSIÓN PLANA
# =============================================================================


def matriz_D() -> np.ndarray:
    """Relación esfuerzo-deformación de tensión plana, en kN/m²."""
    E = D.EC * MPA
    nu = D.NU
    return E / (1 - nu**2) * np.array([[1, nu, 0], [nu, 1, 0], [0, 0, (1 - nu) / 2]])


def rigidez_q4(x: np.ndarray, z: np.ndarray, t: float, Dm: np.ndarray) -> np.ndarray:
    """Rigidez 8×8 de un cuadrilátero de 4 nudos, integrada con 2×2 puntos de Gauss.

    Nudos en orden antihorario: (−1,−1), (1,−1), (1,1), (−1,1) en coordenadas
    naturales. Grados de libertad: (ux, uz) de cada nudo.
    """
    K = np.zeros((8, 8))
    g = 1.0 / math.sqrt(3.0)
    xi_n = np.array([-1, 1, 1, -1])
    eta_n = np.array([-1, -1, 1, 1])
    for xi in (-g, g):
        for eta in (-g, g):
            dN_dxi = 0.25 * xi_n * (1 + eta * eta_n)
            dN_deta = 0.25 * eta_n * (1 + xi * xi_n)
            J = np.array([[dN_dxi @ x, dN_dxi @ z], [dN_deta @ x, dN_deta @ z]])
            dN = np.linalg.solve(J, np.vstack([dN_dxi, dN_deta]))   # [dN/dx; dN/dz]
            B = np.zeros((3, 8))
            B[0, 0::2] = dN[0]
            B[1, 1::2] = dN[1]
            B[2, 0::2] = dN[1]
            B[2, 1::2] = dN[0]
            K += B.T @ Dm @ B * t * np.linalg.det(J)
    return K


# =============================================================================
# 4. CARGAS EN LAS LOSAS
# =============================================================================


def cargas_por_losa(base: list[tuple[float, float, float]],
                    tope: list[float] | None = None) -> list[tuple[float, float, float]]:
    """(H, N, C) de cada losa a partir de (P, V, M) en la base de cada entrepiso.

    Con `tope` (P en el tope de cada entrepiso), N_i = P_tope,i − P_(i+1): la
    losa lleva solo lo que entra en ella, y el peso propio va aparte
    (`pesos_propios`). Sin `tope`, N_i = P_i − P_(i+1): el peso del entrepiso
    queda en la losa de arriba."""
    cargas = []
    n = len(base)
    for i in range(n):
        P, V, M = base[i]
        Pa, Va, Ma = base[i + 1] if i + 1 < n else (0.0, 0.0, 0.0)
        Pt = P if tope is None else tope[i]
        cargas.append((V - Va, Pt - Pa, (M - Ma) + D.ALTURAS[i] * V))
    return cargas


def pesos_propios(base: list[tuple[float, float, float]], tope: list[float]) -> list[float]:
    """W_i = P_i − P_tope,i: el peso propio de cada entrepiso, como carga
    (negativa: hacia abajo), con el factor de la combinación."""
    return [base[i][0] - tope[i] for i in range(len(base))]


def con_peso_propio() -> bool:
    return bool(getattr(D, "AXIAL_CON_PESO_PROPIO", True))


def sismo_entre_base_y_tope() -> bool:
    return bool(getattr(D, "SISMO_ENTRE_BASE_Y_TOPE", True))


def tope_coherente(S_base: float, S_tope: float, S_V: float, h: float) -> bool:
    """¿Cae el tope de ETABS entre |S_base − h·S_V| y S_base + h·S_V? (Lo que
    exige cualquier SRSS o CQC; ver la cabecera.)"""
    b, t, v = abs(S_base), abs(S_tope), abs(S_V)
    tol = 1e-6 * max(b, t, h * v, 1.0)
    return abs(b - h * v) - tol <= t <= b + h * v + tol


def sismo_en_el_entrepiso(S_base: float, S_tope: float, S_V: float, h: float,
                          zeta: float) -> float:
    """La parte de sismo del momento a la altura relativa ζ del entrepiso:
    S(ζ)² = (1 − ζ)·S_base² + ζ·S_tope² − ζ·(1 − ζ)·h²·S_V², exacta para SRSS y
    CQC; si el tope no es coherente con eso, la recta entre base y tope. Con
    el signo de la semidiferencia (el mismo en la base y en el tope)."""
    signo = 1.0 if (S_base if S_base != 0.0 else S_tope) >= 0.0 else -1.0
    b, t, v = abs(S_base), abs(S_tope), abs(S_V)
    if tope_coherente(b, t, v, h):
        r2 = (1.0 - zeta) * b * b + zeta * t * t - zeta * (1.0 - zeta) * h * h * v * v
        return signo * math.sqrt(max(r2, 0.0))
    return signo * ((1.0 - zeta) * b + zeta * t)


def vector_de_cargas(malla: Malla, cargas, pesos: list[float] | None = None) -> np.ndarray:
    """Reparte (H, N, C) de cada losa en los nudos de su plano, y el peso
    propio W de cada entrepiso en sus planos interiores.

    Cada nudo recibe según su ancho tributario (la mitad del elemento a cada
    lado): H, N y W repartidos uniformemente, y C como fuerza vertical lineal
    con el brazo, de resultante nula y momento igual a C. W se reparte entre
    los planos interiores según su altura tributaria (la mitad del elemento de
    arriba y la del de abajo): ni en la base del entrepiso (así la cara de
    encima de la base lleva todo su peso, como el P de la base de ETABS) ni en
    la losa (así la cara de debajo de la losa no lo lleva, como el P del tope).
    """
    n_nudos = malla.nx * len(malla.zs)
    f = np.zeros(2 * n_nudos)
    dx = np.diff(malla.xs)
    ancho = np.zeros(malla.nx)
    ancho[:-1] += dx / 2
    ancho[1:] += dx / 2
    area = ancho * D.ESPESOR
    xc = float((malla.xs * area).sum() / area.sum())
    brazo = malla.xs - xc
    inercia = float((brazo**2 * area).sum())
    for (H, N, C), j in zip(cargas, malla.losas):
        for i in range(malla.nx):
            k = malla.nudo(i, j)
            f[2 * k] += H * area[i] / area.sum()
            f[2 * k + 1] += N * area[i] / area.sum() + C * brazo[i] * area[i] / inercia
    for s, W in enumerate(pesos or []):
        interiores = list(range(malla.bases[s] + 1, malla.losas[s]))
        tributaria = np.array([(malla.zs[j + 1] - malla.zs[j - 1]) / 2 for j in interiores])
        for j, parte in zip(interiores, tributaria / tributaria.sum()):
            for i in range(malla.nx):
                f[2 * malla.nudo(i, j) + 1] += W * parte * area[i] / area.sum()
    return f


# =============================================================================
# 5. RESOLVER Y SACAR FUERZAS DE SECCIÓN
# =============================================================================


@dataclass
class ResultadoMEF:
    malla: Malla
    estados: list                 # [(combinación, paso)]
    encima: np.ndarray            # [estado, plano, (P, V, M)]: cara de ENCIMA del plano
    debajo: np.ndarray            # [estado, plano, (P, V, M)]: cara de DEBAJO del plano
    n_elementos: int


def resolver(fuerzas: dict) -> ResultadoMEF:
    """Resuelve el MEF para todos los estados de carga a la vez.

    `fuerzas` ya en los ejes del programa (`mapear`)."""
    malla = crear_malla()
    nx, nz = malla.nx, len(malla.zs)
    Dm = matriz_D()
    n_gdl = 2 * nx * nz

    # Ensamblaje de la rigidez global
    K = np.zeros((n_gdl, n_gdl))
    elementos = []   # (nudos, gdl, Ke, plano inferior, plano superior)
    for j in range(nz - 1):
        for i in range(nx - 1):
            nudos = [malla.nudo(i, j), malla.nudo(i + 1, j),
                     malla.nudo(i + 1, j + 1), malla.nudo(i, j + 1)]
            x = malla.xs[[i, i + 1, i + 1, i]]
            z = malla.zs[[j, j, j + 1, j + 1]]
            Ke = rigidez_q4(x, z, D.ESPESOR, Dm)
            gdl = np.array([[2 * n, 2 * n + 1] for n in nudos]).ravel()
            K[np.ix_(gdl, gdl)] += Ke
            elementos.append((nudos, gdl, Ke, j, j + 1))

    # Cargas: una columna por estado
    estados = list(fuerzas)
    F = np.zeros((n_gdl, len(estados)))
    for k, clave in enumerate(estados):
        base = [(fuerzas[clave][s]["Bottom"]["P"],
                 fuerzas[clave][s]["Bottom"]["Vx"],
                 fuerzas[clave][s]["Bottom"]["Mx"]) for s in range(len(D.PISOS))]
        if con_peso_propio():
            tope = [fuerzas[clave][s]["Top"]["P"] for s in range(len(D.PISOS))]
            F[:, k] = vector_de_cargas(malla, cargas_por_losa(base, tope), pesos_propios(base, tope))
        else:
            F[:, k] = vector_de_cargas(malla, cargas_por_losa(base))

    # Empotramiento en la base (plano 0): se eliminan esos grados de libertad
    fijos = np.array([[2 * malla.nudo(i, 0), 2 * malla.nudo(i, 0) + 1] for i in range(nx)]).ravel()
    libres = np.setdiff1d(np.arange(n_gdl), fijos)
    U = np.zeros_like(F)
    U[libres] = np.linalg.solve(K[np.ix_(libres, libres)], F[libres])

    # Fuerzas de sección por equilibrio de fuerzas nodales
    xc = D.LW / 2
    encima = np.zeros((len(estados), nz, 3))
    debajo = np.zeros((len(estados), nz, 3))
    for nudos, gdl, Ke, j_inf, j_sup in elementos:
        f = Ke @ U[gdl]                       # fuerzas internas en los 4 nudos, (8, estados)
        # Cara inferior del elemento: sus nudos 0 y 1 están en el plano j_inf.
        # Lo que el material de ARRIBA le transmite a ese plano.
        for loc in (0, 1):
            x = malla.xs[nudos[loc] % nx]
            encima[:, j_inf, 0] -= f[2 * loc + 1]
            encima[:, j_inf, 1] -= f[2 * loc]
            encima[:, j_inf, 2] -= f[2 * loc + 1] * (x - xc)
        # Cara superior: nudos 2 y 3 en el plano j_sup. Lo que llega al
        # material de ABAJO justo bajo ese plano (incluye la carga de la losa).
        for loc in (2, 3):
            x = malla.xs[nudos[loc] % nx]
            debajo[:, j_sup, 0] += f[2 * loc + 1]
            debajo[:, j_sup, 1] += f[2 * loc]
            debajo[:, j_sup, 2] += f[2 * loc + 1] * (x - xc)
    # redondeado a 1e-6 kN (y kN·m): el ruido numérico del sistema lineal
    # (hasta 1e-9 kN) no llega al informe
    return ResultadoMEF(malla, estados, np.round(encima, 6), np.round(debajo, 6), len(elementos))


# =============================================================================
# 6. LA DEMANDA EN CUALQUIER COTA
# =============================================================================


@dataclass
class Estado:
    """Fuerzas de un estado de carga en una cota.

    P negativo en compresión. V y M en el plano del muro (dirección X); Vy y My
    fuera de él (dirección Y). Los momentos, respecto del centroide de la
    sección real y negativos cuando comprimen la cara +X (+Y).
    Vg y VEh: la parte de gravedad y la de sismo (con su signo) de V, de las
    combinaciones y casos de ETABS de SEPARACION_DEL_CORTE. Ve es el cortante
    de diseño en el plano (lo llena diseno_aci.py con la amplificación de ACI
    §18.10.3.3.2, solo de VEh y solo si `con_sismo`); Vu, el mayorado sin
    amplificar. Vy no se usa: el corte en Y no se verifica, porque en Y el muro
    no trabaja como muro (en los tipos II y III las alas son columnas de los
    pórticos especiales); va con My en los signos del sismo.
    """

    combo: str
    paso: str
    piso: int
    z: float
    P: float
    V: float
    M: float
    Vy: float
    My: float
    Vg: float = 0.0
    VEh: float = 0.0
    Ve: float = 0.0
    con_sismo: bool = False

    @property
    def Vu(self) -> float:
        return self.V

    @property
    def nombre(self) -> str:
        return f"{self.combo} {self.paso}".strip()


SIGNOS = [(sp, sm, sy) for sp in (1, -1) for sm in (1, -1) for sy in (1, -1)]

# kN y kN·m: lo que la parte estática del tope de ETABS puede apartarse de la
# estática desde la base por redondeo (ETABS da 0.01; en los muros sin cargas
# entre losas la diferencia no pasa de 0.02). Más que eso: cargas dentro del
# entrepiso (`cargas_dentro_del_entrepiso`).
TOLERANCIA_ESTATICA = 0.10


def cargas_dentro_del_entrepiso(fuerzas: dict, piso: int) -> bool:
    """¿Entra carga dentro del entrepiso en ETABS? La parte estática (la
    gravitacional y la de gravedad de cada envolvente, la media de Max y Min)
    tiene que cumplir V_tope = V_base y M_tope = M_base + V·h, en X y en Y; si
    en alguna combinación no lo hace por más que el redondeo de ETABS, el
    entrepiso tiene cargas entre losas (un nivel intermedio, vigas a media
    altura…). `fuerzas`, en los ejes del programa (`mapear`)."""
    h = D.ALTURAS[piso]
    combos: dict = {}
    for (combo, paso), f in fuerzas.items():
        combos.setdefault(combo, []).append(f[piso])
    for filas in combos.values():
        b = sum(fila5(f["Bottom"]) for f in filas) / len(filas)
        t = sum(fila5(f["Top"]) for f in filas) / len(filas)
        if (abs(t[1] - b[1]) > TOLERANCIA_ESTATICA or abs(t[3] - b[3]) > TOLERANCIA_ESTATICA
                or abs(t[2] - (b[2] + b[1] * h)) > TOLERANCIA_ESTATICA
                or abs(t[4] - (b[4] + b[3] * h)) > TOLERANCIA_ESTATICA):
            return True
    return False


def _signo(s: int) -> str:
    return "+" if s > 0 else "−"


class Demanda:
    """P, V y M de cada estado en cualquier cota, a partir del MEF."""

    def __init__(self, mef: ResultadoMEF, fuerzas: dict, traslado=(0.0, 0.0),
                 corte: "SeparacionDelCorte | None" = None) -> None:
        """`fuerzas` en los ejes del programa; `traslado` = (x_ETABS − x_real,
        y_ETABS − y_real), lo que se mueve el punto de referencia de los
        momentos (ver la cabecera); `corte`, de `leer_separacion_del_corte`:
        de dónde salen Vg y VEh para el cortante de diseño."""
        self.mef = mef
        self.fuerzas = fuerzas
        self.traslado = traslado
        self.corte = corte
        m = mef.malla
        self.z_base = [float(m.zs[j]) for j in m.bases]
        self.z_tope = [float(m.zs[j]) for j in m.losas]
        self.H = self.z_tope[-1]
        self.con_cargas_dentro = [cargas_dentro_del_entrepiso(fuerzas, s) for s in range(len(self.z_base))]

    def planos_interiores(self, piso: int) -> list[float]:
        m = self.mef.malla
        j0, j1 = m.bases[piso], m.losas[piso]
        return [float(m.zs[j]) for j in range(j0 + 1, j1)]

    def en(self, z: float, piso: int) -> list[Estado]:
        """Los estados en la cota z, con el material del entrepiso `piso`.

        En una losa hay dos caras: la de encima (base del entrepiso de arriba)
        y la de debajo (tope del entrepiso de abajo); `piso` decide cuál.
        Dentro del entrepiso, M se interpola entre los valores EXACTOS de los
        planos de nudos (es una recta: entre losas no entra momento ni
        cortante). Con AXIAL_CON_PESO_PROPIO, P va en recta entre la base y el
        tope de ETABS (los dos los devuelve el MEF). Con
        SISMO_ENTRE_BASE_Y_TOPE, la parte de sismo de las envolventes sale de
        la base y el tope de ETABS con `sismo_en_el_entrepiso` (ver la cabecera).
        Si ETABS tiene cargas dentro del entrepiso (`con_cargas_dentro`), van
        además los mismos estados leídos desde el tope de ETABS, con «(tope)»
        en el paso: la demanda es la envolvente de las dos lecturas.
        """
        m = self.mef.malla
        j0, j1 = m.bases[piso], m.losas[piso]
        zs = m.zs[j0:j1 + 1]
        zb, zt = self.z_base[piso], self.z_tope[piso]
        avance = (z - zb) / (zt - zb)
        crudos: dict = {}      # {combinación: {paso: (P, V, M, Vy, My)}}, M respecto de ETABS
        desde_el_tope: dict = {}   # lo mismo, por estática desde el tope de ETABS
        extremos: dict = {}    # {combinación: {paso: (fila de la base, fila del tope)}} de ETABS
        for k, clave in enumerate(self.mef.estados):
            # valores de M en cada plano del entrepiso: cara de encima en la
            # base, cara de debajo en el tope y el valor común en los interiores
            Ms = np.r_[self.mef.encima[k, j0:j1, 2], self.mef.debajo[k, j1, 2]]
            M = float(np.interp(z, zs, Ms))
            P, V = self.mef.encima[k, j0, 0], self.mef.encima[k, j0, 1]
            b = self.fuerzas[clave][piso]["Bottom"]
            t = self.fuerzas[clave][piso]["Top"]
            if con_peso_propio():
                # P en recta entre la base y el tope (con el peso propio)
                P = b["P"] + (t["P"] - b["P"]) * avance
            Vy = b["Vy"]
            My = b["My"] + Vy * (z - zb)
            crudos.setdefault(clave[0], {})[clave[1]] = (float(P), float(V), M, float(Vy), float(My))
            # con cargas dentro del entrepiso, la otra lectura: V, Vy y la
            # estática de M y My desde el tope de ETABS (P, la misma recta)
            desde_el_tope.setdefault(clave[0], {})[clave[1]] = (
                float(P), float(t["Vx"]), float(t["Mx"] - t["Vx"] * (zt - z)), float(t["Vy"]),
                float(t["My"] - t["Vy"] * (zt - z)))
            extremos.setdefault(clave[0], {})[clave[1]] = (
                np.array([b["P"], b["Vx"], b["Mx"], b["Vy"], b["My"]]),
                np.array([t["P"], t["Vx"], t["Mx"], t["Vy"], t["My"]]))

        salida = self._estados(crudos, extremos, piso, z, avance, "Bottom", "")
        if self.con_cargas_dentro[piso]:
            # ETABS tiene cargas dentro del entrepiso (no se sabe dónde): la
            # demanda es la envolvente de las dos lecturas, la estática desde la
            # base y desde el tope (criterio del proyecto)
            salida += self._estados(desde_el_tope, extremos, piso, z, avance, "Top", " (tope)")
        return salida

    def _estados(self, crudos: dict, extremos: dict, piso: int, z: float, avance: float,
                 lado: str, sufijo: str) -> list[Estado]:
        """Los estados de una lectura (`crudos`): la gravedad, de ella; el
        sismo de las envolventes, de la base y el tope de ETABS (`extremos`);
        Vg del corte, de la cara `lado` del entrepiso. `sufijo` va en el paso."""
        zb, zt = self.z_base[piso], self.z_tope[piso]
        salida = []
        dx, dy = self.traslado
        for combo, pasos in crudos.items():
            if set(pasos) != {"Max", "Min"}:
                # sin par Max/Min (la gravitacional): todo es «gravedad»
                for paso, (P, V, M, Vy, My) in pasos.items():
                    salida.append(Estado(combo, (paso + sufijo).strip(), piso, z, P, V, M + P * dx, Vy,
                                         My + P * dy, Vg=V))     # sin sismo: no se amplifica
                continue
            # Envolvente de ETABS: la parte de gravedad es la media de Max y Min
            # y la de sismo, la mitad de la diferencia (con el signo de la fila Max).
            # La de gravedad, del MEF (estática); la de sismo, de la base y el
            # tope de ETABS (o, sin SISMO_ENTRE_BASE_Y_TOPE, de la estática desde
            # la base).
            mx, mn = np.array(pasos["Max"]), np.array(pasos["Min"])
            g, d = (mx + mn) / 2, (mx - mn) / 2
            if sismo_entre_base_y_tope():
                (base_max, tope_max), (base_min, tope_min) = extremos[combo]["Max"], extremos[combo]["Min"]
                d_b, d_t = (base_max - base_min) / 2, (tope_max - tope_min) / 2
                d = d_b + (d_t - d_b) * avance        # P, V y Vy: constantes (base = tope)
                h = zt - zb
                for iM, iV in ((2, 1), (4, 3)):       # Mx con Vx, My con Vy
                    d[iM] = sismo_en_el_entrepiso(d_b[iM], d_t[iM], d_b[iV], h, avance)
            # Para el cortante de diseño (§18.10.3.3.2), Vg y VEh por separado: la
            # combinación de gravedad y el caso de sismo de ETABS de esta
            # combinación (SEPARACION_DEL_CORTE). V es constante en el entrepiso.
            if self.corte is not None and combo in self.corte.partes:
                grav, sis = self.corte.partes[combo]
                Vg = self.corte.gravedad[grav][piso][lado]["Vx"]
                SV = self.corte.sismo[sis][piso]["Bottom"]["Vx"]
            else:
                Vg, SV = g[1], d[1]       # sin separación: la media y la semidiferencia
            if getattr(D, "PERMUTAR_SIGNOS", True):
                signos = SIGNOS
            else:
                signos = [(1, 1, 1), (-1, -1, -1)]          # solo las filas Max y Min
            for sp, sm, sy in signos:
                P = g[0] + sp * d[0]
                V, M = g[1] + sm * d[1], g[2] + sm * d[2]
                Vy, My = g[3] + sy * d[3], g[4] + sy * d[4]
                if len(signos) == 2:
                    paso = "Max" if sp > 0 else "Min"
                else:
                    paso = _signo(sp) + _signo(sm) + _signo(sy)
                paso += sufijo
                # el traslado al centroide real, con el P de ESTE estado
                salida.append(Estado(combo, paso, piso, z, float(P), float(V), float(M + P * dx),
                                     float(Vy), float(My + P * dy), Vg=float(Vg),
                                     VEh=float(sm * SV), con_sismo=True))
        return salida


# =============================================================================
# 6b. EL CORTE: LA PARTE DE GRAVEDAD Y LA DE SISMO POR SEPARADO
# =============================================================================
#
# §18.10.3.3.2 amplifica solo la parte de sismo del cortante, VuEh. Para eso,
# cada combinación con sismo se separa con las combinaciones y casos de ETABS
# (SEPARACION_DEL_CORTE): su parte de gravedad (0.9D, 1.25(L+D)) y su parte de
# sismo (el espectro de diseño), y se comprueba que la combinación de ETABS es
# gravedad ± sismo. Solo para el corte: la flexión, el biaxial, el borde y la
# cota de corte usan las combinaciones de ETABS tal como vienen.

TOLERANCIA_DEL_CORTE = 0.2     # kN: tres redondeos de ETABS (0.01 kN o 0.01 tonf)


@dataclass
class SeparacionDelCorte:
    """partes[combinación] = (gravedad, sismo), los nombres de ETABS.
    gravedad[nombre][piso][ubicación] y sismo[nombre][piso][ubicación], ya en
    los ejes del programa; el sismo, magnitudes con el signo de los ejes (la
    «fila Max» de la parte de sismo: P, V2, V3, M2 y M3 positivos en ETABS).
    comprobacion[combinación] = (mayor diferencia en Vx, mayor en P, V y M en X
    e Y) entre las filas Max y Min de la combinación y gravedad ± sismo, en
    todos los pisos, base y tope."""

    partes: dict
    gravedad: dict
    sismo: dict
    comprobacion: dict


def fila5(r: dict) -> np.ndarray:
    """(P, Vx, Mx, Vy, My) de un registro ya en los ejes del programa."""
    return np.array([r["P"], r["Vx"], r["Mx"], r["Vy"], r["My"]], dtype=float)


def leer_separacion_del_corte(fuerzas: dict) -> SeparacionDelCorte:
    """SEPARACION_DEL_CORTE del archivo de datos, leída de ETABS y comprobada.
    `fuerzas`: las de COMBINACIONES, ya en los ejes del programa (`mapear`)."""
    sep = getattr(D, "SEPARACION_DEL_CORTE", None)
    if not isinstance(sep, dict):
        raise DatoInvalido("SEPARACION_DEL_CORTE es un diccionario {combinación: (gravedad, sismo)} "
                           "con los nombres de ETABS")
    con_sismo = sorted({c for c, paso in fuerzas if paso in ("Max", "Min")})
    for c in con_sismo:
        if c not in sep:
            raise DatoInvalido(f"SEPARACION_DEL_CORTE: falta {c!r}, que lleva sismo (filas Max y Min en "
                               f"ETABS): para amplificar solo el sismo en el corte hacen falta su "
                               f"gravedad y su sismo")
    for c, par in sep.items():
        if c not in con_sismo:
            raise DatoInvalido(f"SEPARACION_DEL_CORTE: {c!r} no es una combinación con sismo de "
                               f"COMBINACIONES (con filas Max y Min en ETABS)")
        if not (isinstance(par, (tuple, list)) and len(par) == 2
                and all(isinstance(x, str) and x != "" for x in par)):
            raise DatoInvalido(f"SEPARACION_DEL_CORTE[{c!r}] es (combinación de gravedad de ETABS, "
                               f"caso de sismo de ETABS); es {par!r}")
    nombres_g = {g for g, _ in sep.values()}
    nombres_s = {s for _, s in sep.values()}
    crudas = leer_fuerzas_etabs(nombres_g | nombres_s)
    faltan = sorted(n for n in nombres_g | nombres_s if not any(k[0] == n for k in crudas))
    if faltan:
        raise DatoInvalido(f"en {D.ARCHIVO_FUERZAS} no está, para el pier {D.MURO}: {faltan} (en "
                           f"SEPARACION_DEL_CORTE van los nombres de ETABS, su «Output Case»)")
    gravedad = {}
    for g in sorted(nombres_g):
        pasos = sorted(k[1] for k in crudas if k[0] == g)
        if pasos != [""]:
            raise DatoInvalido(f"SEPARACION_DEL_CORTE: «{g}» es una parte de gravedad: tiene que ser una "
                               f"combinación de ETABS sin sismo, con una fila por ubicación; trae los "
                               f"pasos {pasos}")
        gravedad[g] = crudas[(g, "")]
    sismo = {}
    for s in sorted(nombres_s):
        pasos = {k[1] for k in crudas if k[0] == s}
        if "Max" not in pasos or not pasos <= {"Max", "Min"}:
            raise DatoInvalido(f"SEPARACION_DEL_CORTE: «{s}» es una parte de sismo: tiene que ser un "
                               f"caso de espectro de ETABS (fila Max); trae los pasos {sorted(pasos)}")
        mx, mn = crudas[(s, "Max")], crudas.get((s, "Min"))
        # un espectro da magnitudes (filas Max, positivas); si trae también Min,
        # la magnitud es la semidiferencia
        sismo[s] = {p: {d: {k: ((mx[p][d][k] - mn[p][d][k]) / 2 if mn else abs(mx[p][d][k]))
                            for k in COMPONENTES} for d in mx[p]} for p in mx}
    gravedad, sismo = mapear(gravedad), mapear(sismo)
    comprobacion = {}
    for c, (g, s) in sep.items():
        en_v = en_todo = 0.0
        for piso in range(len(D.PISOS)):
            for donde in ("Bottom", "Top"):
                G, S = fila5(gravedad[g][piso][donde]), fila5(sismo[s][piso][donde])
                for paso, esperada in (("Max", G + S), ("Min", G - S)):
                    dif = np.abs(fila5(fuerzas[(c, paso)][piso][donde]) - esperada)
                    en_v, en_todo = max(en_v, float(dif[1])), max(en_todo, float(dif.max()))
        if en_v > TOLERANCIA_DEL_CORTE:
            raise DatoInvalido(f"SEPARACION_DEL_CORTE: {c} no es {g} ± {s}: su cortante difiere hasta "
                               f"{en_v:.2f} kN (revisa los nombres o el factor del sismo en ETABS)")
        comprobacion[c] = (en_v, en_todo)
    return SeparacionDelCorte(dict(sep), gravedad, sismo, comprobacion)


# =============================================================================
# 7. CONTRASTE CON ETABS
# =============================================================================


def contraste_con_etabs(mef: ResultadoMEF, fuerzas: dict) -> dict:
    """Diferencias entre el MEF y ETABS.

    Las de ETABS, ya en los ejes del programa (`mapear`), y respecto de su
    propio centro de gravedad, que es como las devuelve el MEF.
    - BASE de cada entrepiso, todos los estados: P, Vx y Mx.
    - TOPE de cada entrepiso, la parte ESTÁTICA: la gravitacional y la parte
      de gravedad de cada envolvente (la media de Max y Min). Tiene que
      cumplir M_tope = M_base + V·h, y con el peso propio repartido P también
      coincide.
    - TOPE, la parte de SISMO de cada envolvente (la semidiferencia): ETABS,
      la estática desde la base y las cotas |S_base − h·S_V| y
      S_base + h·S_V, entre las que tiene que caer ETABS (ver la cabecera); y a
      media altura, la estática, la exacta y la recta.
    La dirección Y (fuera del plano, sin MEF) la contrasta ejecutar.py con la
    demanda ya armada (`contraste_en_y`).
    """
    m = mef.malla
    k_de = {clave: k for k, clave in enumerate(mef.estados)}

    def fila(r: dict) -> np.ndarray:
        return np.array([r["P"], r["Vx"], r["Mx"]])

    base, tope, sismo = [], [], []
    peor = {"P": 0.0, "V": 0.0, "M": 0.0}
    peor_tope = {"P": 0.0, "V": 0.0, "M": 0.0}
    for clave, k in k_de.items():
        for s in range(len(D.PISOS)):
            et, me = fila(fuerzas[clave][s]["Bottom"]), mef.encima[k, m.bases[s]]
            base.append((clave, s, tuple(et), tuple(me)))
            for n, a, b in zip("PVM", et, me):
                peor[n] = max(peor[n], abs(a - b))
    combos: dict = {}
    for clave in mef.estados:
        combos.setdefault(clave[0], []).append(clave[1])
    for combo, pasos in combos.items():
        envolvente = set(pasos) == {"Max", "Min"}
        for s in range(len(D.PISOS)):
            if envolvente:
                kx, kn = k_de[(combo, "Max")], k_de[(combo, "Min")]
                tx, tn = fila(fuerzas[(combo, "Max")][s]["Top"]), fila(fuerzas[(combo, "Min")][s]["Top"])
                bx, bn = fila(fuerzas[(combo, "Max")][s]["Bottom"]), fila(fuerzas[(combo, "Min")][s]["Bottom"])
                mx, mn = mef.debajo[kx, m.losas[s]], mef.debajo[kn, m.losas[s]]
                filas_tope = [(f"{combo} gravedad", (tx + tn) / 2, (mx + mn) / 2)]
                d_base, d_tope, d_estatica = (bx - bn) / 2, (tx - tn) / 2, (mx - mn) / 2
                h = D.ALTURAS[s]
                Sb, St, SV = abs(d_base[2]), abs(d_tope[2]), abs(d_base[1])
                sismo.append({"combo": combo, "piso": s, "base": d_base, "tope": d_tope,
                              "estatica": d_estatica,
                              "inferior": abs(Sb - h * SV), "superior": Sb + h * SV,
                              "coherente": tope_coherente(Sb, St, SV, h),
                              # a media altura: estática desde la base, exacta y recta
                              "medio": (abs(Sb - h / 2 * SV),
                                        abs(sismo_en_el_entrepiso(Sb, St, SV, h, 0.5)),
                                        (Sb + St) / 2)})
            else:
                filas_tope = [(f"{combo} {paso}".strip(), fila(fuerzas[(combo, paso)][s]["Top"]),
                               mef.debajo[k_de[(combo, paso)], m.losas[s]]) for paso in pasos]
            for nombre, et, me in filas_tope:
                tope.append((nombre, s, tuple(et), tuple(me)))
                for n, a, b in zip("PVM", et, me):
                    peor_tope[n] = max(peor_tope[n], abs(a - b))
    return {"base": base, "tope": tope, "sismo": sismo, "peor": peor, "peor_tope": peor_tope}
